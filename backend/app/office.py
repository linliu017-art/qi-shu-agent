"""Bounded office content models. Files never run formulas or fetch resources."""
from io import BytesIO
import json
import math
import re
from typing import Literal
from pydantic import BaseModel, ConfigDict, Field, model_validator
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


class Sheet(BaseModel):
    model_config = ConfigDict(extra='forbid')
    name: str = Field(min_length=1,max_length=31)
    rows: list[list[str | int | float | bool | None]] = Field(min_length=1,max_length=1000)

    @model_validator(mode='after')
    def bounds(self):
        if re.search(r'[\\/*?:\[\]]', self.name) or self.name.startswith("'") or self.name.endswith("'"):
            raise ValueError('工作表名称无效')
        for row in self.rows:
            if not row or len(row)>50:
                raise ValueError('每行须有1至50个单元格')
            for value in row:
                if isinstance(value,str) and (len(value)>5000 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',value)):
                    raise ValueError('单元格内容无效或过长')
                if isinstance(value,float) and not math.isfinite(value):
                    raise ValueError('数字必须有限')
        return self


class Slide(BaseModel):
    model_config = ConfigDict(extra='forbid')
    title: str = Field(min_length=1,max_length=50)
    body: list[str] = Field(min_length=1,max_length=6)

    @model_validator(mode='after')
    def bounds(self):
        if any(not s.strip() or len(s)>100 or '\n' in s for s in self.body):
            raise ValueError('每页1至6条正文，每条1至100字符且不含换行')
        return self


class OfficeDocument(BaseModel):
    model_config = ConfigDict(extra='forbid')
    kind: Literal['xlsx','pptx']
    title: str = Field(min_length=1,max_length=100)
    sheets: list[Sheet] = Field(default_factory=list,max_length=10)
    slides: list[Slide] = Field(default_factory=list,max_length=30)

    @model_validator(mode='after')
    def shape(self):
        strings=[self.title]+[s.name for s in self.sheets]+[v for s in self.slides for v in [s.title,*s.body]]
        if any(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',v) for v in strings):
            raise ValueError('内容含不支持的控制字符')
        if self.kind=='xlsx':
            if not self.sheets or self.slides:
                raise ValueError('Excel仅使用sheets')
            if len({s.name.casefold() for s in self.sheets})!=len(self.sheets):
                raise ValueError('工作表名称不能重复')
        elif not self.slides or self.sheets:
            raise ValueError('PPT仅使用slides')
        if len(self.model_dump_json())>90000:
            raise ValueError('成果超过结构化大小限制')
        return self


def as_text(document):
    lines=['# '+document.title]
    if document.kind=='xlsx':
        for sheet in document.sheets:
            lines.append('\n## '+sheet.name)
            for i,row in enumerate(sheet.rows,1):
                lines.append(f'行{i}: '+ ' | '.join('' if v is None else str(v) for v in row))
    else:
        for i,slide in enumerate(document.slides,1):
            lines.append(f'\n## {i}. {slide.title}\n'+'\n'.join('- '+v for v in slide.body))
    return '\n'.join(lines)


def export_office(document):
    output=BytesIO()
    if document.kind=='xlsx':
        book=Workbook();book.remove(book.active)
        for data in document.sheets:
            sheet=book.create_sheet(data.name)
            for row in data.rows:
                sheet.append(row)
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value,str):
                        cell.data_type='s'  # Formula-like content remains literal text.
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    cell.font=Font(name='Arial',size=11)
            for cell in sheet[1]:
                cell.fill=PatternFill('solid',fgColor='203B63')
                cell.font=Font(name='Arial',bold=True,color='FFFFFF')
            for i in range(1,sheet.max_column+1):
                sheet.column_dimensions[get_column_letter(i)].width=min(48,max(14,max(len(str(sheet.cell(r,i).value or ''))*1.6 for r in range(1,sheet.max_row+1))+2))
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            sheet.page_setup.orientation='landscape'
            sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
            sheet.sheet_properties.pageSetUpPr.fitToPage=True
            sheet.print_title_rows='1:1'
            for row in sheet:
                lines=max(sum(max(1,math.ceil(sum(2 if ord(ch)>127 else 1 for ch in line)/max(1,sheet.column_dimensions[cell.column_letter].width-4))) for line in str(cell.value or '').split('\n')) for cell in row)
                sheet.row_dimensions[row[0].row].height=min(400,max(28,lines*22+12))
        book.save(output)
    else:
        deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
        for index,data in enumerate(document.slides,1):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            title=slide.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(1.1)).text_frame
            title.word_wrap=True;title.text=data.title
            for run in title.paragraphs[0].runs:
                run.font.name='Arial';run.font.size=Pt(28);run.font.bold=True;run.font.color.rgb=RGBColor.from_string('203B63')
            body=slide.shapes.add_textbox(Inches(.75),Inches(1.8),Inches(11.8),Inches(4.95)).text_frame
            body.word_wrap=True
            font_size=18 if sum(math.ceil(len(line)/40) for line in data.body)<=10 else 14
            for i,line in enumerate(data.body):
                p=body.paragraphs[0] if i==0 else body.add_paragraph();p.text='• '+line;p.space_after=Pt(14)
                for run in p.runs: run.font.name='Arial';run.font.size=Pt(font_size);run.font.color.rgb=RGBColor.from_string('26354A')
            footer=slide.shapes.add_textbox(Inches(12),Inches(7),Inches(.6),Inches(.3)).text_frame
            footer.text=str(index)
        deck.save(output)
    return output.getvalue()


class OfficeChange(BaseModel):
    model_config=ConfigDict(extra='forbid')
    path: str = Field(min_length=1,max_length=100)
    original: str | int | float | bool | None
    replacement: str | int | float | bool | None


class OfficeChanges(BaseModel):
    model_config=ConfigDict(extra='forbid')
    edits: list[OfficeChange] = Field(min_length=1,max_length=30)


def apply_office_changes(document,changes):
    """Exact scalar changes only; no arbitrary JSON paths, no silent collateral edits."""
    data=json.loads(json.dumps(document))
    seen=set()
    for change in changes.edits:
        if change.path in seen: raise ValueError('修改路径重复')
        seen.add(change.path)
        if not re.fullmatch(r'title|sheets/\d+/name|sheets/\d+/rows/\d+/\d+|slides/\d+/title|slides/\d+/body/\d+',change.path):
            raise ValueError('不支持的编辑位置')
        parts=change.path.split('/');current=data
        try:
            for part in parts[:-1]:current=current[int(part)] if isinstance(current,list) else current[part]
            key=int(parts[-1]) if isinstance(current,list) else parts[-1]
            if type(current[key]) is not type(change.original) or current[key]!=change.original:
                raise ValueError('原值已变化，请重新读取')
            current[key]=change.replacement
        except (KeyError,IndexError,TypeError):raise ValueError('编辑位置不存在') from None
    return OfficeDocument.model_validate(data)
