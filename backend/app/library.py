"""Immutable material revisions and task-bound provenance; owner filtered locally."""
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
from datetime import datetime, timezone
from uuid import uuid4
from hashlib import sha256
import json
import re
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from . import store

MAX_BYTES=10*1024*1024


def init(db):
    db.execute('CREATE TABLE IF NOT EXISTS materials (id TEXT PRIMARY KEY, owner_id TEXT NOT NULL, name TEXT NOT NULL, kind TEXT NOT NULL, sha256 TEXT NOT NULL, parent_id TEXT, created_at TEXT NOT NULL, content BLOB NOT NULL, text TEXT NOT NULL)')
    db.execute('CREATE TABLE IF NOT EXISTS task_sources (task_id TEXT NOT NULL, source_id TEXT NOT NULL, PRIMARY KEY(task_id,source_id))')
    db.execute('CREATE TABLE IF NOT EXISTS office_artifacts (task_id TEXT PRIMARY KEY, document TEXT NOT NULL)')
    db.execute('CREATE TABLE IF NOT EXISTS native_office_artifacts (task_id TEXT PRIMARY KEY, kind TEXT NOT NULL, sha256 TEXT NOT NULL, content BLOB NOT NULL)')
    db.execute('CREATE TABLE IF NOT EXISTS native_office_sessions (id TEXT PRIMARY KEY, owner_id TEXT NOT NULL, base_id TEXT NOT NULL, kind TEXT NOT NULL, snapshot BLOB NOT NULL, snapshot_hash TEXT NOT NULL, status TEXT NOT NULL, result_id TEXT, result_hash TEXT, recovery BLOB, created_at TEXT NOT NULL)')
    if 'review_details' not in {row['name'] for row in db.execute('PRAGMA table_info(native_office_sessions)')}:
        db.execute("ALTER TABLE native_office_sessions ADD COLUMN review_details TEXT")


def inspect_file(name, raw, *, allow_empty=False):
    name=re.sub(r'[\x00-\x1f\x7f]', '', name.replace('\\','/').split('/')[-1])[:160]
    if not name or not raw or len(raw)>MAX_BYTES:
        raise ValueError('文件不能为空，且不能超过10MB')
    kind=name.rsplit('.',1)[-1].lower()
    if kind not in ('txt','md','docx','xlsx','pptx'):
        raise ValueError('仅支持TXT、Markdown、DOCX、XLSX、PPTX')
    parts=[]
    if kind in ('txt','md'):
        text=raw.decode('utf-8-sig')
        if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',text):
            raise ValueError('文件不是有效文本')
        parts=[text]
    else:
        if not raw.startswith(b'PK\x03\x04'):
            raise ValueError('文件真实类型与扩展名不符')
        with ZipFile(BytesIO(raw)) as archive:
            infos=archive.infolist();names={i.filename for i in infos}
            if len(infos)>2000 or sum(i.file_size for i in infos)>40*1024*1024 or len(names)!=len(infos):
                raise ValueError('压缩包条目或解包大小超过限制')
            expected={'docx':'word/document.xml','xlsx':'xl/workbook.xml','pptx':'ppt/presentation.xml'}[kind]
            if expected not in names or '[Content_Types].xml' not in names:
                raise ValueError('Office文件内容类型无效')
            for info in infos:
                if info.flag_bits & 1 or info.filename.startswith('/') or '\\' in info.filename or '..' in PurePosixPath(info.filename).parts or 'vba' in info.filename.lower():
                    raise ValueError('拒绝加密、宏或不安全路径')
                if info.file_size>max(1,info.compress_size)*300:
                    raise ValueError('异常压缩比')
                if info.filename.endswith(('.xml','.rels')):
                    data=archive.read(info)
                    if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
                        raise ValueError('拒绝XML实体定义')
                    if info.filename.endswith('.rels'):
                        root=ET.fromstring(data)
                        if any(r.attrib.get('TargetMode')=='External' for r in root):
                            raise ValueError('不支持带外部链接的Office文件，请移除链接后上传')
            content_types=archive.read('[Content_Types].xml').lower()
            if b'macroenabled' in content_types:
                raise ValueError('拒绝含宏的文件')
        if kind=='docx':
            doc=Document(BytesIO(raw))
            # Body order is significant; retain paragraph/table positions.
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            pnum=tnum=0
            for block in doc.element.body:
                if block.tag.endswith('}p'):
                    pnum+=1
                    content=Paragraph(block,doc).text
                    if content.strip(): parts.append(f'[段落{pnum}] '+content)
                elif block.tag.endswith('}tbl'):
                    tnum+=1
                    for i,row in enumerate(Table(block,doc).rows,1):
                        values=[c.text for c in row.cells]
                        if any(value.strip() for value in values): parts.append(f'[表格{tnum} 行{i}] '+' | '.join(values))
        elif kind=='xlsx':
            book=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                if len(book.worksheets)>20: raise ValueError('最多支持20个工作表')
                for sheet in book:
                    # Some valid XLSX writers omit the optional dimension tag;
                    # openpyxl then reports None rather than a numeric maximum.
                    if (sheet.max_row is not None and sheet.max_row>10000) or (sheet.max_column is not None and sheet.max_column>100):
                        raise ValueError('表格超过10000行或100列')
                    for row_number,row in enumerate(sheet,1):
                        if row_number>10000 or len(row)>100:raise ValueError('表格超过10000行或100列')
                        values=[]
                        for cell in row:
                            if cell.value is not None and str(cell.value).strip():
                                value=str(cell.value)
                                values.append(f'{cell.coordinate}='+ ('[公式，未计算] ' if cell.data_type=='f' else '')+value)
                        if values: parts.append(f'[工作表 {sheet.title}] '+' | '.join(values))
            finally: book.close()
        else:
            deck=Presentation(BytesIO(raw))
            if len(deck.slides)>100: raise ValueError('最多支持100页幻灯片')
            for i,slide in enumerate(deck.slides,1):
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        if shape.text.strip(): parts.append(f'[幻灯片{i}] '+shape.text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            values=[c.text for c in row.cells]
                            if any(value.strip() for value in values): parts.append(f'[幻灯片{i} 表格] '+' | '.join(values))
    text='\n'.join(parts)
    if not text.strip() and not allow_empty: raise ValueError('文件没有可提取的文字，不支持图片文字识别')
    if len(text)>200000: raise ValueError('提取文本超过20万字符，请拆分文件')
    return name,kind,text


def metadata(row):
    return {k:row[k] for k in ('id','name','kind','sha256','parent_id','created_at')}


def upload(name,raw,parent_id=None):
    name,kind,text=inspect_file(name,raw)
    tid=uuid4().hex
    with store.connection() as db:
        if parent_id and not db.execute("SELECT id FROM materials WHERE id=? AND owner_id='local-user'",(parent_id,)).fetchone():
            raise LookupError('父资料版本不存在')
        db.execute('INSERT INTO materials VALUES (?,?,?,?,?,?,?,?,?)',(tid,'local-user',name,kind,sha256(raw).hexdigest(),parent_id,datetime.now(timezone.utc).isoformat(),raw,text))
    return metadata(get(tid))


def get(tid):
    with store.connection() as db:
        r=db.execute("SELECT * FROM materials WHERE id=? AND owner_id='local-user'",(tid,)).fetchone()
        return dict(r) if r else None


def list_materials(offset=0,limit=30):
    with store.connection() as db:
        rows=db.execute("SELECT id,name,kind,sha256,parent_id,created_at FROM materials WHERE owner_id='local-user' ORDER BY created_at DESC LIMIT ? OFFSET ?",(limit,offset)).fetchall()
        return [dict(r) for r in rows]


def bind(db,tid,ids):
    if len(ids)>10 or len(set(ids))!=len(ids): raise ValueError('最多关联10份不同资料')
    for sid in ids:
        if not db.execute("SELECT id FROM materials WHERE id=? AND owner_id='local-user'",(sid,)).fetchone():
            raise ValueError('资料不存在或无权限')
        db.execute('INSERT INTO task_sources VALUES (?,?)',(tid,sid))


def inherit(db,parent_id,tid):
    from . import review
    review.inherit(db,parent_id,tid)
    db.execute('INSERT INTO task_sources SELECT ?,source_id FROM task_sources WHERE task_id=?',(tid,parent_id))


def sources(tid):
    with store.connection() as db:
        rows=db.execute("SELECT m.id,m.name,m.kind,m.sha256,m.parent_id,m.created_at FROM materials m JOIN task_sources s ON s.source_id=m.id JOIN tasks t ON t.id=s.task_id WHERE t.id=? AND t.owner_id='local-user' AND m.owner_id='local-user' ORDER BY m.created_at",(tid,)).fetchall()
        return [dict(r) for r in rows]


def read_source(tid,sid):
    if sid not in {m['id'] for m in sources(tid)}: raise ValueError('该资料未关联当前任务或无权限')
    return get(sid)


def office(tid):
    with store.connection() as db:
        r=db.execute("SELECT o.document FROM office_artifacts o JOIN tasks t ON t.id=o.task_id WHERE t.id=? AND t.owner_id='local-user'",(tid,)).fetchone()
        return json.loads(r[0]) if r else None


def finish_office(tid,document,last_correction,review_items=None):
    from .office import as_text
    with store.connection() as db:
        db.execute('BEGIN IMMEDIATE')
        cursor=db.execute("UPDATE tasks SET status='completed',progress='办公成果已保存',result=? WHERE id=? AND owner_id='local-user' AND status='running' AND NOT EXISTS (SELECT 1 FROM corrections WHERE task_id=? AND id>?)",(as_text(document),tid,tid,last_correction))
        if cursor.rowcount!=1: return False
        db.execute('INSERT INTO office_artifacts VALUES (?,?)',(tid,document.model_dump_json()))
        from . import review
        review.write_items(db,tid,review_items)
        return True


def edit(tid,expected,content=None,document=None):
    from .office import as_text
    new_id=uuid4().hex
    with store.connection() as db:
        db.execute('BEGIN IMMEDIATE')
        family=store.versions(tid)
        if not family or tid!=expected or family[-1]['id']!=expected:
            raise ValueError('版本已变化，请刷新最新版本后编辑')
        row=db.execute("SELECT * FROM tasks WHERE id=? AND owner_id='local-user' AND status='completed'",(tid,)).fetchone()
        if not row or db.execute('SELECT 1 FROM task_context WHERE parent_id=?',(tid,)).fetchone():
            raise ValueError('该版本已有后续任务，请刷新后编辑')
        if db.execute('SELECT 1 FROM native_office_artifacts WHERE task_id=?',(tid,)).fetchone():
            raise ValueError('复杂办公成果须通过原文件编辑，不能使用简化编辑')
        existing=db.execute('SELECT document FROM office_artifacts WHERE task_id=?',(tid,)).fetchone()
        if bool(existing)!=bool(document): raise ValueError('请使用与原成果类型对应的编辑接口')
        if existing and json.loads(existing[0])['kind']!=document.kind: raise ValueError('不能通过编辑更换成果格式')
        result=as_text(document) if document else content
        db.execute('INSERT INTO tasks VALUES (?,?,?,?,?,?,?,?)',(new_id,'local-user','用户直接编辑成果，以本版本为准',row['material'],'completed','编辑已保存为新版本',result,datetime.now(timezone.utc).isoformat()))
        db.execute('INSERT INTO task_context VALUES (?,?,?)',(new_id,tid,'completed'))
        inherit(db,tid,new_id)
        if document: db.execute('INSERT INTO office_artifacts VALUES (?,?)',(new_id,document.model_dump_json()))
    return store.get(new_id)


def native_office(tid):
    """Read the immutable native Office bytes of an owned task version."""
    with store.connection() as db:
        row=db.execute("SELECT n.kind,n.sha256,n.content FROM native_office_artifacts n JOIN tasks t ON t.id=n.task_id WHERE t.id=? AND t.owner_id='local-user' AND t.status='completed'",(tid,)).fetchone()
        return dict(row) if row else None


def native_office_info(tid):
    """Get native Office metadata without loading file bytes into memory."""
    with store.connection() as db:
        row=db.execute("SELECT n.kind,n.sha256 FROM native_office_artifacts n JOIN tasks t ON t.id=n.task_id WHERE t.id=? AND t.owner_id='local-user' AND t.status='completed'",(tid,)).fetchone()
        return dict(row) if row else None


def import_native_office(name, raw):
    """Create an immutable original and its first editable task atomically."""
    name,kind,text=inspect_file(name,raw,allow_empty=True)
    if kind not in ('docx','xlsx','pptx'):
        raise ValueError('仅支持Word、Excel和PPT原文件')
    sid,tid=uuid4().hex,uuid4().hex
    now=datetime.now(timezone.utc).isoformat();digest=sha256(raw).hexdigest()
    summary='# 原文件编辑成果\n\n'+(text[:95000] if text.strip() else '文件没有可提取的正文，内容请在原文件编辑器中查看。')
    with store.connection() as db:
        db.execute('BEGIN IMMEDIATE')
        db.execute('INSERT INTO materials VALUES (?,?,?,?,?,?,?,?,?)',(sid,'local-user',name,kind,digest,None,now,raw,text))
        db.execute('INSERT INTO tasks VALUES (?,?,?,?,?,?,?,?)',(tid,'local-user','编辑原文件：'+name,'','completed','原文件已保存',summary,now))
        db.execute('INSERT INTO task_sources VALUES (?,?)',(tid,sid))
        db.execute('INSERT INTO native_office_artifacts VALUES (?,?,?,?)',(tid,kind,digest,raw))
    return store.get(tid)


def _save_native_tx(db, tid, expected, kind, raw, text):
    """Save a native version inside the caller's existing write transaction."""
    family=store.versions(tid)
    if not family or tid!=expected or family[-1]['id']!=expected:
        raise ValueError('版本已变化，请刷新最新版本后编辑')
    row=db.execute("SELECT * FROM tasks WHERE id=? AND owner_id='local-user' AND status='completed'",(tid,)).fetchone()
    if not row or db.execute('SELECT 1 FROM task_context WHERE parent_id=?',(tid,)).fetchone():
        raise ValueError('该版本已有后续任务，请刷新后编辑')
    native=db.execute('SELECT kind FROM native_office_artifacts WHERE task_id=?',(tid,)).fetchone()
    structured=db.execute('SELECT document FROM office_artifacts WHERE task_id=?',(tid,)).fetchone()
    previous_kind=native['kind'] if native else json.loads(structured['document'])['kind'] if structured else 'docx'
    if kind!=previous_kind:
        raise ValueError('不能通过编辑更换成果格式')
    digest=sha256(raw).hexdigest();new_id=uuid4().hex
    summary='# 原文件编辑成果\n\n'+(text[:95000] if text.strip() else '文件没有可提取的正文，内容请在原文件编辑器中查看。')
    db.execute('INSERT INTO tasks VALUES (?,?,?,?,?,?,?,?)',(new_id,'local-user','用户直接编辑原文件，以本版本为准',row['material'],'completed','原文件已保存为新版本',summary,datetime.now(timezone.utc).isoformat()))
    db.execute('INSERT INTO task_context VALUES (?,?,?)',(new_id,tid,'completed'))
    inherit(db,tid,new_id)
    db.execute('INSERT INTO native_office_artifacts VALUES (?,?,?,?)',(new_id,kind,digest,raw))
    return new_id


def save_native_version(tid, expected, name, raw):
    """Validate and atomically preserve editor-returned bytes as a new version.

    This is a storage primitive, not an editor endpoint. The integration must
    validate editor callbacks, sessions and user authorization separately.
    """
    _,kind,text=inspect_file(name,raw,allow_empty=True)
    if kind not in ('docx','xlsx','pptx'):
        raise ValueError('仅支持Word、Excel和PPT原文件')
    with store.connection() as db:
        db.execute('BEGIN IMMEDIATE')
        new_id=_save_native_tx(db,tid,expected,kind,raw,text)
    return store.get(new_id)


def prepare_native_editor_file(tid):
    """Build a local editor input without modifying any existing task version."""
    task=store.get(tid)
    if not task or task['status']!='completed':
        raise ValueError('仅支持已完成的成果')
    native=native_office(tid)
    if native:
        return native
    structured=office(tid)
    if structured:
        from .office import OfficeDocument,export_office
        document=OfficeDocument.model_validate(structured)
        kind=document.kind
        raw=export_office(document)
    else:
        from .word_export import export_docx
        kind='docx'
        raw=export_docx(task['result'])
    # These are local generated files, but still apply the same format/size checks.
    inspect_file('result.'+kind,raw,allow_empty=True)
    return {'kind':kind,'sha256':sha256(raw).hexdigest(),'content':raw}
