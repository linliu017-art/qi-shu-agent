"""Local Markdown-to-DOCX export. No HTML execution or external fetches."""
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from markdown_it import MarkdownIt


def inline(paragraph, token):
    bold = italic = strike = False
    for child in token.children or []:
        if child.type == 'strong_open': bold = True
        elif child.type == 'strong_close': bold = False
        elif child.type == 'em_open': italic = True
        elif child.type == 'em_close': italic = False
        elif child.type == 's_open': strike = True
        elif child.type == 's_close': strike = False
        elif child.type in ('softbreak', 'hardbreak'): paragraph.add_run().add_break()
        elif child.type in ('text', 'code_inline', 'image'):
            run = paragraph.add_run(child.content)
            run.bold, run.italic, run.font.strike = bold, italic, strike
            if child.type == 'code_inline': run.font.name = 'Courier New'
        elif child.type == 'link_close':
            pass  # Visible label preserved; no external network requests.


def export_docx(markdown):
    document = Document()
    document.core_properties.author = '启枢'
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.8)
    style = document.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
    for name in ['Title']+[f'Heading {n}' for n in range(1,7)]:
        heading = document.styles[name]
        heading.font.name = 'Arial'
        heading.font.color.rgb = RGBColor(0,0,0)
        heading.font.size = Pt(20 if name=='Title' else max(12,18-int(name[-1])))
        heading.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Songti SC')
        for border in heading.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    style.paragraph_format.space_after = Pt(8)
    tokens = MarkdownIt('commonmark', {'html': False}).enable('table').enable('strikethrough').parse(markdown)
    paragraph = None
    list_styles = []
    table = row = cell = None
    header = False
    for token in tokens:
        kind = token.type
        if kind == 'heading_open':
            paragraph = document.add_paragraph(style='Title' if token.tag=='h1' else f'Heading {min(int(token.tag[1]),6)}')
        elif kind == 'bullet_list_open': list_styles.append('List Bullet')
        elif kind == 'ordered_list_open': list_styles.append('List Number')
        elif kind in ('bullet_list_close','ordered_list_close'): list_styles.pop()
        elif kind == 'paragraph_open':
            paragraph = document.add_paragraph(style=list_styles[-1] if list_styles else 'Normal')
        elif kind == 'inline':
            if paragraph is None: paragraph = document.add_paragraph()
            inline(paragraph,token)
        elif kind in ('fence','code_block'):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(token.content.rstrip('\n'))
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif kind == 'table_open':
            table = document.add_table(rows=0,cols=0)
            table.style = 'Table Grid'
        elif kind == 'thead_open': header = True
        elif kind == 'thead_close': header = False
        elif kind == 'tr_open':
            row = table.add_row()
            cell = 0
            if header:
                repeat = OxmlElement('w:tblHeader')
                row._tr.get_or_add_trPr().append(repeat)
        elif kind in ('th_open','td_open'):
            if cell >= len(table.columns): table.add_column(Inches(1))
            paragraph = row.cells[cell].paragraphs[0]
            cell += 1
        elif kind == 'table_close':
            count = len(table.columns)
            weights = [max(7,min(30,max(len(r.cells[i].text) for r in table.rows))) for i in range(count)]
            table.autofit = False
            for i,col in enumerate(table.columns):
                width = Inches(6.6*weights[i]/sum(weights))
                col.width = width
                for r in table.rows: r.cells[i].width = width
            table = None
    output = BytesIO()
    document.save(output)
    return output.getvalue()
