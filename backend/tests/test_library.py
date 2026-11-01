from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from concurrent.futures import ThreadPoolExecutor
import json
import pytest
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook,load_workbook
from pptx import Presentation
from app import store,library,agent
from app.main import app
from app.office import OfficeDocument,export_office


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('AGENT_DB_PATH',str(tmp_path/'data.sqlite'))
    monkeypatch.delenv('MODEL_API_KEY',raising=False)
    with TestClient(app) as client: yield client


def sample(kind):
    if kind in ('txt','md'):return '虚构访谈：预算未定，负责人张三。'.encode()
    out=BytesIO()
    if kind=='docx':
        doc=Document();doc.add_heading('客户访谈',0);doc.add_paragraph('预算未定，负责人张三。');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='范围';table.cell(0,1).text='产品部';doc.save(out)
    elif kind=='xlsx':
        book=Workbook();book.active.title='需求';book.active.append(['负责人','预算']);book.active.append(['张三','=1+1']);book.save(out)
    else:
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='需求概要';slide.placeholders[1].text='预算未定，负责人张三。';deck.save(out)
    return out.getvalue()


def done(sources=None):
    task=store.create('整理','虚构资料',sources or [])
    store.update(task['id'],'running','测试')
    assert store.finish(task['id'],'原成果',0)
    return task['id']


@pytest.mark.parametrize('kind',['txt','md','docx','xlsx','pptx'])
def test_upload_roundtrip_and_locations(client,kind):
    raw=sample(kind)
    r=client.post('/api/v1/materials',files={'file':(f'../../客户.{kind}',raw)})
    assert r.status_code==201,r.text
    data=r.json();assert data['name']==f'客户.{kind}'
    preview=client.get('/api/v1/materials/'+data['id']).json()
    assert '张三' in preview['text']
    if kind=='xlsx':assert '[公式，未计算] =1+1' in preview['text'] and 'A2' in preview['text']
    assert client.get('/api/v1/materials/'+data['id']+'/download').content==raw
    assert client.get('/api/v1/materials/'+data['id'],params={'offset':999999}).status_code==400


@pytest.mark.parametrize('name,raw',[('x.docx',b'not-a-zip'),('x.exe',b'plain'),('x.txt',b'\x00bad'),('x.md',b'\xff'),('x.txt',b'a'*200001)])
def test_bad_files(client,name,raw):
    assert client.post('/api/v1/materials',files={'file':(name,raw)}).status_code==400


@pytest.mark.parametrize('member,payload',[('../evil.xml',b'x'),('word/vbaProject.bin',b'x'),('word/large.xml',b'x'*200000),('word/_rels/document.xml.rels',b'<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>')])
def test_hostile_office(client,member,payload):
    data=BytesIO()
    with ZipFile(data,'w',ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','<Types/>');z.writestr('word/document.xml','<document/>');z.writestr(member,payload)
    assert client.post('/api/v1/materials',files={'file':('bad.docx',data.getvalue())}).status_code==400


def test_material_versions_sources_and_owner(client):
    first=library.upload('v1.txt',b'first')
    second=library.upload('v2.txt',b'second',first['id'])
    tid=done([first['id']]);child=store.create_followup(tid,'修改')['id']
    assert library.sources(child)[0]['id']==first['id']
    assert library.read_source(child,first['id'])['text']=='first'
    with pytest.raises(ValueError):library.read_source(child,second['id'])
    assert library.get(second['id'])['parent_id']==first['id']
    with store.connection() as db:db.execute("UPDATE materials SET owner_id='other' WHERE id=?",(second['id'],))
    assert client.get('/api/v1/materials/'+second['id']).status_code==404
    before=len(store.list_tasks())
    assert client.post('/api/v1/tasks',json={'goal':'test','source_ids':[second['id']]}).status_code==400
    assert len(store.list_tasks())==before
    assert client.post('/api/v1/materials',data={'parent_id':second['id']},files={'file':('new.txt',b'new')}).status_code==404


def test_text_edit_conflict_atomic_and_source_inheritance(client):
    src=library.upload('source.txt',b'fact');tid=done([src['id']])
    def edit():
        try:return library.edit(tid,tid,content='新稿')['id']
        except ValueError:return None
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:edit(),range(2)))
    assert sum(r is not None for r in results)==1
    assert store.get(tid)['result']=='原成果'
    new=next(r for r in results if r)
    assert library.sources(new)==library.sources(tid)
    assert client.post(f'/api/v1/tasks/{tid}/edit',json={'content':'过期','expected_version_id':tid}).status_code==409
    assert client.get(f'/api/v1/tasks/{new}/export/md').content.decode()=='新稿'
    store.init();assert library.sources(new)[0]['id']==src['id']


def office_data(kind):
    if kind=='xlsx':return {'kind':kind,'title':'需求清单','sheets':[{'name':'需求','rows':[['事项','状态'],['预算','未定'],['公式文本','=1+1']]}]}
    return {'kind':kind,'title':'需求汇报','slides':[{'title':'已知事实','body':['预算未定','负责人张三']},{'title':'待确认','body':['上线日期未知']} ]}


@pytest.mark.parametrize('kind',['xlsx','pptx'])
def test_office_edit_export_restore_and_no_model(client,kind):
    tid=store.create('生成办公成果','虚构资料')['id'];store.update(tid,'running','测试')
    document=OfficeDocument.model_validate(office_data(kind));assert library.finish_office(tid,document,0)
    data=client.get(f'/api/v1/tasks/{tid}/export/{kind}');assert data.status_code==200
    if kind=='xlsx':
        book=load_workbook(BytesIO(data.content));assert book['需求']['B3'].value=='=1+1' and book['需求']['B3'].data_type=='s'
    else:
        deck=Presentation(BytesIO(data.content));assert len(deck.slides)==2
        assert any('预算未定' in s.text for s in deck.slides[0].shapes if s.has_text_frame)
    document.title='用户编辑后'
    r=client.post(f'/api/v1/tasks/{tid}/office/edit',json={'expected_version_id':tid,'document':document.model_dump()})
    assert r.status_code==201;r_id=r.json()['id']
    assert library.office(r_id)['title']=='用户编辑后'
    assert library.office(tid)['title']!='用户编辑后'
    restored=store.restore(r_id,tid)
    assert library.office(restored['id'])==library.office(tid)
    assert store.task_usage(r_id)['input_tokens']==0
    assert client.post(f'/api/v1/tasks/{r_id}/edit',json={'expected_version_id':r_id,'content':'失去结构'}).status_code==409


def test_invalid_office_and_correction_guard(client):
    data=office_data('xlsx');data['sheets'].append(data['sheets'][0])
    with pytest.raises(ValueError):OfficeDocument.model_validate(data)
    data=office_data('pptx');data['slides'][0]['body']=['x'*101]
    with pytest.raises(ValueError):OfficeDocument.model_validate(data)
    tid=store.create('office','')['id'];store.update(tid,'running','test');store.add_correction(tid,'新增要求')
    assert not library.finish_office(tid,OfficeDocument.model_validate(office_data('xlsx')),0)
    assert library.office(tid) is None


def test_agent_reads_bound_file_guide_and_office(client,monkeypatch):
    source=library.upload('访谈.txt','预算未定，负责人张三'.encode())
    task=store.create('整理需求为Excel','',[source['id']]);monkeypatch.setenv('MODEL_API_KEY','test-key');monkeypatch.setenv('MODEL_BASE_URL','https://example.test');monkeypatch.setenv('MODEL_NAME','test-model')
    calls=[('read_business_guide',{'guide_id':'requirements'}),('read_source',{'source_id':source['id']}),('save_office',office_data('xlsx'))]
    seen=[]
    class Client:
        def __init__(self,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def post(self,*args,**kwargs):
            seen.append(kwargs['json']['messages'][:]);name,arguments=calls[len(seen)-1]
            class Response:
                def raise_for_status(self):pass
                def json(self):return {'choices':[{'message':{'role':'assistant','tool_calls':[{'id':str(len(seen)),'type':'function','function':{'name':name,'arguments':json.dumps(arguments)}}]}}],'usage':{'prompt_tokens':10,'completion_tokens':10}}
            return Response()
    monkeypatch.setattr(agent.httpx,'Client',Client);agent.run(task['id'])
    assert store.get(task['id'])['status']=='completed'
    assert library.office(task['id'])['kind']=='xlsx'
    assert '张三' in seen[-1][-1]['content']


def test_office_exact_patch_preserves_other_cells_and_rejects_stale():
    from app.office import OfficeChanges,apply_office_changes
    original=office_data('xlsx')
    edits=OfficeChanges.model_validate({'edits':[{'path':'sheets/0/rows/1/1','original':'未定','replacement':'待客户确认'}]})
    result=apply_office_changes(original,edits)
    expected=json.loads(json.dumps(original));expected['sheets'][0]['rows'][1][1]='待客户确认'
    assert result==OfficeDocument.model_validate(expected)
    assert original['sheets'][0]['rows'][1][1]=='未定'
    with pytest.raises(ValueError):apply_office_changes(result.model_dump(),edits)
    with pytest.raises(ValueError):apply_office_changes(original,OfficeChanges.model_validate({'edits':[{'path':'sheets/0/rows/-1/0','original':None,'replacement':'x'}]}))


def test_upload_size_and_api_office_errors(client):
    assert client.post('/api/v1/materials',files={'file':('large.txt',b'a'*(library.MAX_BYTES+1))}).status_code==413
    tid=done()
    assert client.get(f'/api/v1/tasks/{tid}/export/xlsx').status_code==409
    assert client.get(f'/api/v1/tasks/{tid}/export/exe').status_code==400
    assert client.get('/api/v1/materials',params={'limit':1000}).status_code==422
    assert len(client.get('/api/v1/business-guides').json()['guides'])==7


def test_office_empty_or_image_only_is_not_mistaken_for_text(client):
    from PIL import Image
    from pptx.util import Inches
    image=BytesIO();Image.new('RGB',(8,8),'white').save(image,format='PNG');pixels=image.getvalue()
    word=Document();word.add_picture(BytesIO(pixels));word.add_table(rows=1,cols=2)
    word_bytes=BytesIO();word.save(word_bytes)
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(1));slide.shapes.add_table(1,2,Inches(1),Inches(2),Inches(2),Inches(1));slide.shapes.add_picture(BytesIO(pixels),Inches(1),Inches(3))
    ppt_bytes=BytesIO();deck.save(ppt_bytes)
    excel=Workbook();excel.active['A1']='   ';excel_bytes=BytesIO();excel.save(excel_bytes)
    for name,raw in [('image.docx',word_bytes.getvalue()),('image.pptx',ppt_bytes.getvalue()),('blank.xlsx',excel_bytes.getvalue())]:
        response=client.post('/api/v1/materials',files={'file':(name,raw)})
        assert response.status_code==400, (name,response.text)
        assert '没有可提取的文字' in response.json()['error']['message']


def test_mixed_office_keeps_real_text_locations_and_original(client):
    from PIL import Image
    from pptx.util import Inches
    image=BytesIO();Image.new('RGB',(8,8),'white').save(image,format='PNG');pixels=image.getvalue()
    doc=Document();doc.add_picture(BytesIO(pixels));doc.add_paragraph('真实文字');table=doc.add_table(rows=2,cols=2);table.cell(1,0).text='表中资料';word_bytes=BytesIO();doc.save(word_bytes)
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(BytesIO(pixels),Inches(1),Inches(1));slide.shapes.add_textbox(Inches(1),Inches(2),Inches(3),Inches(1)).text='幻灯片文字';ppt_bytes=BytesIO();deck.save(ppt_bytes)
    excel=Workbook();excel.active['A1']='  ';excel.active['A2']=0;excel_bytes=BytesIO();excel.save(excel_bytes)
    for name,raw,expected in [('mixed.docx',word_bytes.getvalue(),'[段落2] 真实文字'),('mixed.pptx',ppt_bytes.getvalue(),'[幻灯片1] 幻灯片文字'),('zero.xlsx',excel_bytes.getvalue(),'A2=0')]:
        response=client.post('/api/v1/materials',files={'file':(name,raw)})
        assert response.status_code==201,(name,response.text)
        ident=response.json()['id'];assert expected in client.get('/api/v1/materials/'+ident).json()['text']
        assert client.get('/api/v1/materials/'+ident+'/download').content==raw
