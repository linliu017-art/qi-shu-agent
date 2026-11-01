from concurrent.futures import ThreadPoolExecutor
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from pptx import Presentation
from pptx.util import Inches
from app import library, store
from app.office import OfficeDocument


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setenv('AGENT_DB_PATH',str(tmp_path/'native.sqlite'))
    store.init()


def make_result(kind):
    task=store.create('虚构成果','虚构资料')
    store.update(task['id'],'running','测试')
    if kind=='docx':
        assert store.finish(task['id'],'# 虚构成果',0)
    else:
        data={'kind':kind,'title':'虚构成果'}
        if kind=='xlsx':data['sheets']=[{'name':'数据','rows':[['标题'],['原值']]}]
        else:data['slides']=[{'title':'原页面','body':['原文字']}]
        assert library.finish_office(task['id'],OfficeDocument.model_validate(data),0)
    return task['id']


def office_file(kind):
    picture=BytesIO();Image.new('RGB',(12,12),'blue').save(picture,format='PNG');image=picture.getvalue();out=BytesIO()
    if kind=='docx':
        doc=Document();doc.sections[0].header.paragraphs[0].text='页眉';doc.add_picture(BytesIO(image));doc.save(out)
    elif kind=='xlsx':
        book=Workbook();sheet=book.active;sheet.append(['数值']);sheet.append([10]);sheet['A3']='=SUM(A2:A2)';sheet.merge_cells('A5:B5');chart=BarChart();chart.add_data(Reference(sheet,min_col=1,min_row=1,max_row=2),titles_from_data=True);sheet.add_chart(chart,'D1');book.save(out)
    else:
        deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_picture(BytesIO(image),Inches(1),Inches(1));slide.notes_slide.notes_text_frame.text='演讲者备注';deck.save(out)
    return out.getvalue()


def test_excel_without_optional_dimension_is_valid():
    import re
    original=office_file('xlsx');out=BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(out,'w',ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data=source.read(entry.filename)
            if entry.filename=='xl/worksheets/sheet1.xml':data=re.sub(rb'<dimension\b[^>]*/>',b'',data)
            target.writestr(entry.filename,data)
    _,kind,text=library.inspect_file('test.xlsx',out.getvalue(),allow_empty=True)
    assert kind=='xlsx' and 'A2=10' in text and '[公式，未计算]' in text


@pytest.mark.parametrize('kind',['docx','xlsx','pptx'])
def test_native_version_preserves_exact_bytes_across_restart_and_restore(db,kind):
    parent=make_result(kind);raw=office_file(kind)
    new=library.save_native_version(parent,parent,'复杂成果.'+kind,raw);tid=new['id']
    stored=library.native_office(tid)
    assert stored['kind']==kind and stored['content']==raw and stored['sha256']==sha256(raw).hexdigest()
    assert library.native_office(parent) is None
    with pytest.raises(ValueError,match='简化编辑'):
        library.edit(tid,tid,content='破坏结构')
    store.init();assert library.native_office(tid)['content']==raw
    restored=store.restore(tid,tid)
    assert library.native_office(restored['id'])['content']==raw
    if kind=='xlsx':
        sheet=load_workbook(BytesIO(raw))['Sheet'];assert sheet['A3'].value=='=SUM(A2:A2)' and len(sheet._charts)==1 and len(sheet.merged_cells.ranges)==1
    elif kind=='pptx':
        deck=Presentation(BytesIO(raw));assert '演讲者备注' in deck.slides[0].notes_slide.notes_text_frame.text
    else:
        assert Document(BytesIO(raw)).sections[0].header.paragraphs[0].text=='页眉'


def test_native_rejects_stale_mismatched_and_unsafe_input_without_writing(db):
    parent=make_result('xlsx');raw=office_file('xlsx');before=len(store.list_tasks())
    for filename,content in [('bad.xlsx',b'not-a-zip'),('wrong.docx',office_file('docx'))]:
        with pytest.raises(ValueError):library.save_native_version(parent,parent,filename,content)
    assert len(store.list_tasks())==before
    data=BytesIO()
    with ZipFile(data,'w',ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','<Types/>');z.writestr('xl/workbook.xml','<workbook/>');z.writestr('../unsafe.xml','bad')
    with pytest.raises(ValueError):library.save_native_version(parent,parent,'unsafe.xlsx',data.getvalue())
    new=library.save_native_version(parent,parent,'good.xlsx',raw)
    with pytest.raises(ValueError,match='版本已变化'):
        library.save_native_version(parent,parent,'good.xlsx',raw)
    assert library.native_office(new['id'])['content']==raw


def test_native_concurrent_saves_create_only_one_new_version(db):
    parent=make_result('pptx');raw=office_file('pptx')
    def save():
        try:return library.save_native_version(parent,parent,'x.pptx',raw)['id']
        except ValueError:return None
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:save(),range(2)))
    assert sum(result is not None for result in results)==1
    assert len(store.versions(parent))==2


def test_native_office_enforces_owner_and_restores_sources(db):
    source=library.upload('虚构来源.txt','预算待确认'.encode())
    root=store.create('虚构任务','',[source['id']]);store.update(root['id'],'running','测试');assert store.finish(root['id'],'# 成果',0)
    raw=office_file('docx');new=library.save_native_version(root['id'],root['id'],'成果.docx',raw)
    assert [m['id'] for m in library.sources(new['id'])]==[source['id']]
    with store.connection() as conn:conn.execute("UPDATE tasks SET owner_id='other' WHERE id=?",(new['id'],))
    assert library.native_office(new['id']) is None
    with pytest.raises(ValueError):library.save_native_version(new['id'],new['id'],'成果.docx',raw)


@pytest.mark.parametrize('kind',['docx','xlsx','pptx'])
def test_editor_input_preserves_native_bytes_and_does_not_create_version(db,kind):
    parent=make_result(kind)
    before=len(store.versions(parent));initial=library.prepare_native_editor_file(parent)
    assert initial['kind']==kind and initial['sha256']==sha256(initial['content']).hexdigest()
    assert len(store.versions(parent))==before
    raw=office_file(kind);new=library.save_native_version(parent,parent,'result.'+kind,raw)
    later=library.prepare_native_editor_file(new['id'])
    assert later['content']==raw and later['sha256']==sha256(raw).hexdigest()
    assert len(store.versions(new['id']))==before+1


def test_editor_input_rejects_unfinished_or_other_owner(db):
    task=store.create('未完成任务','');assert task['status']=='queued'
    with pytest.raises(ValueError):library.prepare_native_editor_file(task['id'])
    tid=make_result('docx')
    with store.connection() as conn:conn.execute("UPDATE tasks SET owner_id='other' WHERE id=?",(tid,))
    with pytest.raises(ValueError):library.prepare_native_editor_file(tid)
