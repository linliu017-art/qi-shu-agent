"""Real, bounded synthetic model test. Preserves reports, no user database edits."""
import os,sys,json,time
from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO
from datetime import datetime,timezone
ROOT=Path(__file__).resolve().parents[2];sys.path.insert(0,str(ROOT/'backend'))
from app import store,library,agent
from app.office import OfficeDocument,export_office
from docx import Document
assert os.environ.get('EVAL_APPROVED')=='yes','须有用户评测授权'
rows=[]
trace=[]
OriginalClient=agent.httpx.Client
class ObservedClient(OriginalClient):
 def post(self,*args,**kwargs):
  response=super().post(*args,**kwargs)
  if response.is_success:
   for choice in response.json().get('choices',[]):
    trace.extend(call.get('function',{}).get('name') for call in choice.get('message',{}).get('tool_calls',[]))
  return response
agent.httpx.Client=ObservedClient
with TemporaryDirectory(prefix='qi-shu-office-eval-') as tmp:
 os.environ['AGENT_DB_PATH']=str(Path(tmp)/'data.sqlite');store.init()
 doc=Document();doc.add_heading('虚构访谈',0);doc.add_paragraph('第一批产品部试用；负责人张三。预算与上线日期尚未确定。');doc.add_paragraph('需要集中查看任务、负责人和延期状态。');out=BytesIO();doc.save(out)
 src=library.upload('虚构访谈.docx',out.getvalue())
 cases=[('xlsx','根据关联访谈调用requirements指引，生成Excel需求清单，列事项、资料事实或待确认、来源位置。不要编造预算和日期。'),('pptx','根据关联访谈调用proposal指引，生成2页PPT，分别列资料明确事实和待确认事项；每页两条简短正文，每条不超过40字符。引用资料位置，预算日期未知，不编造。')]
 for kind,goal in cases:
  trace.clear()
  task=store.create(goal,'',[src['id']]);started=time.monotonic();agent.run(task['id']);result=store.get(task['id']);office=library.office(task['id'])
  row={'tool_calls':trace[:],'kind':kind,'status':result['status'],'progress':result['progress'],'seconds':round(time.monotonic()-started,2),'usage':store.task_usage(task['id']),'sources':library.sources(task['id']),'document':office,'result':result['result']};rows.append(row)
  print(kind,result['status'],flush=True)
  if office:
   artifact=ROOT/'data/evaluation'/f'office-{kind}-latest.{kind}';artifact.write_bytes(export_office(OfficeDocument.model_validate(office)))
   trace.clear()
   followup=store.create_followup(task['id'],'只把负责人张三改成李四，其余数据、页数、文字和来源保持不变；保留办公格式。')
   started=time.monotonic();agent.run(followup['id']);updated=store.get(followup['id']);new=library.office(followup['id'])
   expected=json.loads(json.dumps(office,ensure_ascii=False).replace('张三','李四'))
   rows.append({'tool_calls':trace[:],'kind':kind+'-edit','status':updated['status'],'seconds':round(time.monotonic()-started,2),'usage':store.task_usage(followup['id']),'exact_match':new==expected,'document':new,'original_unchanged':library.office(task['id'])==office})
   print(kind+'-edit',updated['status'],new==expected,flush=True)
 path=ROOT/'data/evaluation'/('office-workflow-'+datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')+'.json')
 path.write_text(json.dumps(rows,ensure_ascii=False,indent=2));print(path)
